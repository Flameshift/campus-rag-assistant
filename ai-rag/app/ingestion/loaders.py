"""各文件格式的 Loader，统一输出 DocumentBlock。"""

import csv
import io
import json
import re
import tempfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from langchain_community.document_loaders import PyPDFLoader

from .blocks import DocumentBlock, normalize_document_text

SUPPORTED_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".log", ".yaml", ".yml", ".xml"}


def _decode_bytes(content: bytes) -> str:
    """优先按 UTF-8 解码，兼容常见中文 CSV/TXT 编码。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _heading_info(text: str) -> tuple[int, str] | None:
    """识别 Markdown、章、节、条和编号标题。"""
    markdown = re.match(r"^(#{1,6})\s+(.+?)\s*#*$", text)
    if markdown:
        return len(markdown.group(1)), markdown.group(2).strip()

    chinese = re.match(r"^(第[一二三四五六七八九十百零〇0-9]+[章节编部分条])\s*(.*)$", text)
    if chinese:
        marker, title = chinese.groups()
        if "章" in marker or "编" in marker or "部分" in marker:
            level = 1
        elif "节" in marker:
            level = 2
        else:
            level = 3
        return level, normalize_document_text(f"{marker} {title}")

    # 单级数字和中文编号通常是条款下的列表项，不能升级为章节标题。
    numbered = re.match(r"^((?:\d+\.)+\d+)\s+(.+)$", text)
    if numbered:
        return 4, text.strip()
    return None


def _blocks_from_text(
    text: str,
    source: str,
    page_start: int | None = None,
    page_end: int | None = None,
    extra_metadata: dict[str, Any] | None = None,
) -> list[DocumentBlock]:
    """按空行和标题行把文本拆为结构块。"""
    blocks: list[DocumentBlock] = []
    buffer: list[str] = []
    order = 0

    def flush() -> None:
        nonlocal order
        content = normalize_document_text("\n".join(buffer))
        buffer.clear()
        if not content:
            return
        content_lines = content.splitlines()
        list_pattern = re.compile(r"^(?:[-*•]\s+|\d+[、.．)]\s*|[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）)")
        block_type = "list" if all(list_pattern.match(line.lstrip()) for line in content_lines) else "paragraph"
        blocks.append(
            DocumentBlock(
                content=content,
                block_type=block_type,
                order_index=order,
                page_start=page_start,
                page_end=page_end,
                source=source,
                metadata=dict(extra_metadata or {}),
            )
        )
        order += 1

    for raw_line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            flush()
            continue

        heading = _heading_info(line)
        if heading:
            flush()
            level, title = heading
            blocks.append(
                DocumentBlock(
                    content=title,
                    block_type="heading",
                    heading_level=level,
                    order_index=order,
                    page_start=page_start,
                    page_end=page_end,
                    source=source,
                    metadata=dict(extra_metadata or {}),
                )
            )
            order += 1
            continue

        buffer.append(line)

    flush()
    return blocks


def _flatten_record(value: Any, prefix: str = "") -> Any:
    """把嵌套 JSON 转成保留字段路径的普通结构。"""
    if isinstance(value, dict):
        return {f"{prefix}.{key}" if prefix else key: _flatten_record(item, f"{prefix}.{key}" if prefix else key) for key, item in value.items()}
    if isinstance(value, list):
        if all(not isinstance(item, (dict, list)) for item in value):
            return value
        return [_flatten_record(item, prefix) for item in value]
    return value


def _blocks_from_csv(content: bytes, source: str) -> list[DocumentBlock]:
    text = _decode_bytes(content)
    reader = csv.DictReader(io.StringIO(text))
    headers = [header.strip() for header in (reader.fieldnames or []) if header is not None]
    blocks: list[DocumentBlock] = []
    for row_index, row in enumerate(reader):
        lines = []
        for header in headers:
            value = row.get(header, "")
            lines.append(f"{header}: {'' if value is None else value}")
        blocks.append(
            DocumentBlock(
                content="\n".join(lines),
                block_type="record",
                order_index=row_index,
                source=source,
                metadata={
                    "record_index": row_index,
                    "record_type": "csv",
                    "record_collection": source,
                },
            )
        )
    return blocks


def _json_items(value: Any) -> tuple[list[Any], str]:
    if isinstance(value, list):
        return value, "$"
    if isinstance(value, dict):
        for key, item in value.items():
            if isinstance(item, list) and item and all(isinstance(entry, dict) for entry in item):
                return item, key
        return [value], "$"
    return [value], "$"


def _blocks_from_json(content: bytes, source: str) -> list[DocumentBlock]:
    parsed = json.loads(_decode_bytes(content))
    blocks: list[DocumentBlock] = []
    items, collection = _json_items(parsed)
    for index, item in enumerate(items):
        flattened = _flatten_record(item)
        content_text = json.dumps(flattened, ensure_ascii=False, indent=2)
        blocks.append(
            DocumentBlock(
                content=content_text,
                block_type="record",
                order_index=index,
                source=source,
                metadata={
                    "record_index": index,
                    "record_type": "json",
                    "record_collection": collection,
                },
            )
        )
    return blocks


def _blocks_from_pdf(filename: str, content: bytes) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    with tempfile.TemporaryDirectory() as tmp_dir:
        path = Path(tmp_dir) / filename
        path.write_bytes(content)
        for page_number, document in enumerate(PyPDFLoader(str(path)).load(), start=1):
            page_blocks = _blocks_from_text(
                document.page_content,
                filename,
                page_start=page_number,
                page_end=page_number,
                extra_metadata={"page": page_number},
            )
            for block in page_blocks:
                block.order_index = len(blocks)
                blocks.append(block)
    return blocks


def _docx_heading_level(style_name: str | None) -> int:
    if not style_name:
        return 0
    value = style_name.strip().lower()
    if "heading" not in value and "标题" not in value:
        return 0
    match = re.search(r"(\d+)", value)
    return int(match.group(1)) if match else 1


def _table_to_markdown(rows: list[list[str]]) -> str:
    clean_rows = [[cell.replace("\n", " ").strip() for cell in row] for row in rows if row]
    if not clean_rows:
        return ""
    width = max(len(row) for row in clean_rows)
    clean_rows = [row + [""] * (width - len(row)) for row in clean_rows]
    header = clean_rows[0]
    separator = ["---"] * width
    body = clean_rows[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _blocks_from_docx(filename: str, content: bytes) -> list[DocumentBlock]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    with tempfile.TemporaryDirectory() as tmp_dir:
        path = Path(tmp_dir) / filename
        path.write_bytes(content)
        document = docx.Document(str(path))

    blocks: list[DocumentBlock] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _docx_heading_level(paragraph.style.name if paragraph.style else "")
            block_type = "heading" if level else ("list" if "list" in (paragraph.style.name or "").lower() else "paragraph")
            blocks.append(
                DocumentBlock(
                    content=text,
                    block_type=block_type,
                    heading_level=level,
                    order_index=len(blocks),
                    source=filename,
                )
            )
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            table_text = _table_to_markdown(rows)
            if table_text:
                blocks.append(
                    DocumentBlock(
                        content=table_text,
                        block_type="table",
                        order_index=len(blocks),
                        source=filename,
                        metadata={"row_count": len(rows), "table_id": len(blocks)},
                    )
                )
    return blocks


class _HTMLTextParser(HTMLParser):
    """轻量 HTML 解析器，无额外依赖地保留标题、段落和列表。"""

    def __init__(self) -> None:
        super().__init__()
        self.blocks: list[tuple[str, int, str]] = []
        self._tag = ""
        self._level = 0
        self._buffer: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        tag = tag.lower()
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._flush()
            self._tag = tag
            self._level = int(tag[1])
        elif tag in {"p", "li"}:
            self._flush()
            self._tag = tag
            self._level = 0

    def handle_data(self, data: str) -> None:
        if self._tag:
            value = data.strip()
            if value:
                self._buffer.append(value)

    def handle_endtag(self, tag: str) -> None:
        if self._tag and tag.lower() == self._tag:
            self._flush()

    def _flush(self) -> None:
        text = normalize_document_text(" ".join(self._buffer))
        if text:
            block_type = "heading" if self._level else ("list" if self._tag == "li" else "paragraph")
            self.blocks.append((block_type, self._level, text))
        self._tag = ""
        self._level = 0
        self._buffer.clear()


def _blocks_from_html(content: bytes, source: str) -> list[DocumentBlock]:
    parser = _HTMLTextParser()
    parser.feed(_decode_bytes(content))
    parser.close()
    return [
        DocumentBlock(
            content=text,
            block_type=block_type,
            heading_level=level,
            order_index=index,
            source=source,
        )
        for index, (block_type, level, text) in enumerate(parser.blocks)
    ]


def _blocks_from_xlsx(filename: str, content: bytes) -> list[DocumentBlock]:
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("解析 xlsx 需要安装 openpyxl") from exc

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)

    blocks: list[DocumentBlock] = []
    for sheet in workbook.worksheets:
        rows = [[("" if value is None else str(value)) for value in row] for row in sheet.iter_rows(values_only=True)]
        table_text = _table_to_markdown(rows)
        if table_text:
            blocks.append(
                DocumentBlock(
                    content=table_text,
                    block_type="table",
                    order_index=len(blocks),
                    source=filename,
                    metadata={"sheet_name": sheet.title, "row_count": len(rows)},
                )
            )
    return blocks


def load_blocks_from_file(filename: str, content: bytes) -> list[DocumentBlock]:
    """根据文件格式解析文档，并统一返回 DocumentBlock。"""
    suffix = Path(filename).suffix.lower()
    safe_name = Path(filename).name or f"upload{suffix}"

    if suffix == ".pdf":
        return _blocks_from_pdf(safe_name, content)
    if suffix == ".docx":
        return _blocks_from_docx(safe_name, content)
    if suffix in {".html", ".htm"}:
        return _blocks_from_html(content, safe_name)
    if suffix == ".csv":
        return _blocks_from_csv(content, safe_name)
    if suffix == ".json":
        return _blocks_from_json(content, safe_name)
    if suffix == ".xlsx":
        return _blocks_from_xlsx(safe_name, content)
    if suffix in SUPPORTED_TEXT_SUFFIXES:
        return _blocks_from_text(_decode_bytes(content), safe_name)

    # 未知扩展名先尝试文本解析，二进制文件解析失败后再报告不支持。
    try:
        text = _decode_bytes(content)
        if "\x00" not in text:
            return _blocks_from_text(text, safe_name)
    except Exception:
        pass
    raise ValueError(f"不支持的文件类型: {filename}")


def blocks_from_text(text: str, source: str = "text.txt") -> list[DocumentBlock]:
    """供 add_text 或测试使用的纯文本入口。"""
    return _blocks_from_text(text, source)
